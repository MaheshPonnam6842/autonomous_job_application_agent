"""Build a single-column, ATS-safe .docx from the structured optimized resume.

ATS-safety choices: one column, standard section headings, no tables / text boxes
/ images / headers-footers, plain hyphen/bulleted lists, a common font. These are
exactly the things resume parsers handle reliably.
"""

from __future__ import annotations

from pathlib import Path


def build_ats_docx(struct: dict, out_path: str | Path) -> str:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(40)
        section.left_margin = section.right_margin = Pt(54)

    def heading(title: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(12)

    header = struct.get("header") or []
    if header:
        p = doc.add_paragraph()
        run = p.add_run(header[0])
        run.bold = True
        run.font.size = Pt(16)
        for extra in header[1:]:
            sub = doc.add_paragraph().add_run(extra)
            sub.font.size = Pt(10)
    links = struct.get("links") or []
    if links:
        doc.add_paragraph().add_run(" | ".join(links)).font.size = Pt(10)

    if struct.get("summary"):
        heading("Professional Summary")
        doc.add_paragraph(struct["summary"])

    if struct.get("skills"):
        heading("Skills")
        doc.add_paragraph(", ".join(struct["skills"]))

    if struct.get("experience"):
        heading("Experience")
        for e in struct["experience"]:
            cp = doc.add_paragraph()
            cp.paragraph_format.space_before = Pt(6)
            cp.add_run(e.get("company", "")).bold = True
            role, dates = e.get("role", ""), e.get("dates", "")
            if role or dates:
                rp = doc.add_paragraph()
                rp.add_run(role).italic = True
                if dates:
                    rp.add_run(f"   |   {dates}")
            for b in e.get("bullets", []):
                text = b["text"] if isinstance(b, dict) else b
                doc.add_paragraph(text, style="List Bullet")

    if struct.get("projects"):
        heading("Projects")
        for p in struct["projects"]:
            doc.add_paragraph().add_run(p.get("name", "")).bold = True
            for b in p.get("bullets", []):
                doc.add_paragraph(b, style="List Bullet")

    if struct.get("education"):
        heading("Education")
        for line in struct["education"]:
            doc.add_paragraph(line)

    out = str(out_path)
    doc.save(out)
    return out
